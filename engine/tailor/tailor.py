"""Tailor: keyword reweighting of the ATS-safe master resumes. Never fabricates.

Operations (reorder and rephrase-within-bounds only):
  1. Summary lead: prepend the matched target title ("Chief Data Officer | ...").
  2. Target roles line: matched title moves first in the pipe list.
  3. Skills / Core Competencies: category paragraphs reordered by posting keyword
     hits; within a single-run category, comma-separated items reordered the same way.

Command: python -m engine.main tailor <application_id>
"""
from __future__ import annotations

import datetime
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document

from ..config import artifacts_dir, profile
from ..db import log_event, tx

TEMPLATE_DIR = Path(__file__).resolve().parent

SKILLS_HEADING = {"data_ai": "Skills", "wellness": "Core Competencies"}
NEXT_HEADING = {"data_ai": "Experience", "wellness": "Wellness Leadership Experience"}

# Sample highlights used in cover notes. Replace with the candidate's own
# quantified highlights (these ship as neutral examples, not real claims).
HIGHLIGHTS = {
    "data_ai": [
        "15+ years building and scaling enterprise data organizations, with documented value creation through AI transformation",
        "Stood up an enterprise AI Center of Excellence: production ML models, measured ROI, and a global team on a multi-million-dollar portfolio",
        "Board-ready: presented transformation roadmaps and ROI outcomes to executive committees at large enterprises",
    ],
    "wellness": [
        "Founded an integrative wellness practice with hundreds of documented client outcomes",
        "Created a structured nervous-system recovery framework taught through books, masterclasses, and corporate programs",
        "Two decades of enterprise leadership designing behavior-change programs adopted by thousands of employees",
    ],
}


def _set_para_text(para, new_text: str) -> None:
    """Replace text in place, preserving paragraph + first-run character style."""
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ""


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _find_paras(doc, heading: str):
    """Return list of paragraphs between `heading` and the next Heading-styled para."""
    out, capturing = [], False
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style else ""
        if _norm(p.text) == heading:
            capturing = True
            continue
        if capturing:
            if style.startswith("Heading") and _norm(p.text):
                break
            if _norm(p.text):
                out.append(p)
    return out


def _hit_count(text: str, keywords: list[str]) -> int:
    low = text.lower()
    return sum(1 for k in keywords if k.lower() in low)


def _reorder_items(text: str, keywords: list[str]) -> str:
    """Reorder comma-separated items in a single-run skills line by keyword hits."""
    m = re.match(r"^(?P<label>[^:]*:\s*)(?P<items>.*)$", text, re.DOTALL)
    if not m:
        return text
    items = [i.strip() for i in m.group("items").split(",") if i.strip()]
    scored = sorted(items, key=lambda i: -_hit_count(i, keywords))
    return m.group("label") + ", ".join(scored)


def tailor_resume(app: dict, posting: dict, matched_title: str, keywords: list[str],
                  out_path: Path) -> list[str]:
    """Produce the tailored variant. Returns a human-readable diff summary."""
    lane = app["lane"]
    tpl = TEMPLATE_DIR / profile(lane)["resume_template"]
    doc = Document(str(tpl))
    changes: list[str] = []

    # 1 + 2: summary lead + target roles line -------------------------------
    for p in doc.paragraphs:
        t = _norm(p.text)
        if t.startswith("Enterprise Data & AI Transformation executive") or \
           t.startswith("Wellness executive, author, and practitioner-researcher"):
            new = f"{matched_title} | {t}"
            _set_para_text(p, new)
            changes.append(f"summary: prepended target title '{matched_title}'")
        elif t.startswith("Target roles:"):
            head, _, rest = t.partition(".")
            roles = [r.strip() for r in head.replace("Target roles:", "").split("|")]
            roles = [matched_title] + [r for r in roles if r and r.lower() != matched_title.lower()]
            deduped = list(dict.fromkeys(roles))
            new = "Target roles: " + " | ".join(deduped) + "." + rest
            _set_para_text(p, new)
            changes.append(f"target line: '{matched_title}' moved first")

    # 3: skills / core competencies reorder ---------------------------------
    skills = _find_paras(doc, SKILLS_HEADING[lane])
    if skills:
        scored = [(_hit_count(p.text, keywords), p) for p in skills]
        ordered = [p for _, p in sorted(scored, key=lambda x: -x[0])]
        if [p.text for p in ordered] != [p.text for p in skills]:
            anchor = skills[0]._p
            for p in ordered:
                anchor.addprevious(p._p)  # move element, preserving styles
            top_hits = scored and max(s for s, _ in scored)
            changes.append(
                f"skills: reordered {len(skills)} categories by posting keyword hits "
                f"(top category matched {top_hits} keywords)"
            )
        for p in skills:
            if len(p.runs) == 1 and "," in p.text and len(p.text) > 60:
                new = _reorder_items(_norm(p.text), keywords)
                if new != _norm(p.text):
                    _set_para_text(p, new)
                    changes.append(f"skills: reordered items within '{p.text.split(':')[0]}:'")

    doc.save(str(out_path))

    # PDF twin when LibreOffice is available
    soffice = shutil.which("soffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_path.parent), str(out_path)],
            capture_output=True, timeout=120,
        )
        changes.append("pdf twin exported")
    else:
        changes.append("pdf export skipped (LibreOffice not found)")
    return changes


def keywords_for(posting: dict, prof: dict) -> list[str]:
    """Profile keywords + meaningful title tokens, longest first for stable hits."""
    kws = list(prof["fit_signals"]["keywords"])
    title_tokens = [t for t in re.split(r"[^A-Za-z&]+", posting.get("title") or "") if len(t) > 3]
    return sorted(set(kws + title_tokens), key=len, reverse=True)


def run(application_id: int) -> None:
    from .cover_note import write_note

    with tx() as conn:
        app = conn.execute(
            "SELECT a.*, p.company, p.title, p.description, p.url, s.rationale "
            "FROM applications a JOIN postings p ON p.id = a.posting_id "
            "LEFT JOIN scores s ON s.posting_id = a.posting_id AND s.lane = a.lane "
            "WHERE a.id = %s", (application_id,)).fetchone()
        if not app:
            raise SystemExit(f"application {application_id} not found")
        if app["stage"] not in ("targeted",):
            print(f"note: application stage is '{app['stage']}' (tailor normally runs from 'targeted')")

        prof = profile(app["lane"])
        matched = None
        try:
            matched = app["rationale"]["matched_title"]
        except Exception:
            pass
        if not matched or matched.startswith("("):  # waived/watchlist pseudo-title
            def _n(s: str) -> str:
                return re.sub(r"[^a-z0-9]+", " ", s.lower().replace("&", " and ")).strip()
            norm_title = _n(app["title"] or "")
            matched = next((t for t in prof["must_have"]["titles_include"] if _n(t) in norm_title),
                           {"data_ai": "VP, Data & Analytics", "wellness": "VP of Wellbeing"}[app["lane"]])

        today = datetime.date.today()
        slug = re.sub(r"[^A-Za-z0-9]+", "_", f"{app['company']}_{app['title']}")[:60]
        out_docx = artifacts_dir() / f"{slug}_{today:%Y-%m-%d}.docx"
        kws = keywords_for({"title": app["title"]}, prof)

        changes = tailor_resume(dict(app), dict(app), matched, kws, out_docx)
        note_path = write_note(dict(app), prof, application_id)

        conn.execute(
            """UPDATE applications SET stage = 'tailored', resume_path = %s, note_path = %s,
               updated_at = now() WHERE id = %s""",
            (str(out_docx), str(note_path), application_id))
        log_event(conn, application_id, "note", {"changes": changes})
        conn.commit()

    print(f"tailored: {out_docx}")
    pdf = out_docx.with_suffix(".pdf")
    if pdf.exists():
        print(f"          {pdf}")
    print(f"cover note: {note_path}")
    print("\nchanges vs master:")
    for c in changes:
        print(f"  - {c}")
    print(f"\nnext: python -m engine.main review {application_id}")
